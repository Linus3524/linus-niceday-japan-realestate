"""Import NPA's 2023 prefecture counts. Requires openpyxl for read-only extraction.
Run with the bundled Python runtime. No API key. Validate before replacing snapshot.
"""
import io
import json
import re
import urllib.request
from datetime import datetime, timezone
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
YEAR = 2023  # Must match crimePrefectureMeta.fiscalYear; never silently mix years.
BASE = "https://www.npa.go.jp/toukei/soubunkan/R05/"
SOURCES = [BASE + "excel/R05_003.xlsx", BASE + "excel/R05_005.xlsx"]
books = [openpyxl.load_workbook(io.BytesIO(urllib.request.urlopen(url).read()), read_only=True, data_only=True) for url in SOURCES]

def compact(value):
    return re.sub(r"\s+", "", str(value or ""))

# Use the existing prefecture list to exclude police regional subtotals and Hokkaido branches.
snapshot = (ROOT / "src/data/crimePrefectureSnapshot.ts").read_text()
names = re.findall(r'"prefecture": "([^"]+)"', snapshot)
assert len(names) == 47 and len(set(names)) == 47
short = {name if name == "北海道" else name[:-1]: name for name in names}

def counts(sheet):
    assert sheet.cell(5, 3).value == "認知件数", sheet.title
    assert str(YEAR) in str(sheet.cell(18, 2).value), sheet.title
    result = {}
    # Rows before 確認用 contain the prefectures; do not ingest hidden check columns.
    for row in sheet.iter_rows(min_row=20, values_only=True):
        name = compact(row[1])
        if name == "確認用":
            break
        if name not in short:
            continue
        value = row[2]
        assert isinstance(value, (int, float)) and value >= 0 and int(value) == value, (sheet.title, name, value)
        assert short[name] not in result
        result[short[name]] = int(value)
    assert len(result) == 47, (sheet.title, len(result))
    assert sum(result.values()) == sheet.cell(18, 3).value, (sheet.title, "national sum")
    return result

groups = [("C", "竊盜犯罪"), ("B", "粗暴犯罪"), ("A", "凶惡犯罪"), ("D", "詐欺等知能犯罪"), ("E", "風俗犯罪"), ("F", "其他刑法犯罪")]
translations = {
    "A-a": "殺人", "A-b": "強盜", "A-c": "縱火", "A-d": "不同意性交等（含致死傷）",
    "B-a": "攜械聚集", "B-b": "暴行", "B-c": "傷害", "B-d": "脅迫", "B-e": "恐嚇取財",
    "D-a": "詐欺", "D-b": "侵占", "D-c": "偽造", "D-d": "貪瀆", "D-e": "斡旋牟利", "D-f": "背信",
    "E-a": "賭博", "E-b": "猥褻相關犯罪", "E-ｃ": "性姿態偷拍等",
    "C-a-(1)": "住宅空巢竊盜", "C-a-(2)": "住宅夜間潛入", "C-a-(3)": "住宅在宅時潛入",
    "C-a-(4)": "ATM 侵入竊盜", "C-a-(5)": "保險箱竊盜", "C-a-(6)": "旅館侵入竊盜",
    "C-a-(7)": "公家機關侵入竊盜", "C-a-(8)": "學校侵入竊盜", "C-a-(9)": "醫院侵入竊盜",
    "C-a-(10)": "加油站侵入竊盜", "C-a-(11)": "辦公室侵入竊盜", "C-a-(12)": "店鋪侵入竊盜",
    "C-a-(13)": "工廠侵入竊盜", "C-a-(14)": "更衣室侵入竊盜", "C-a-(15)": "倉庫侵入竊盜", "C-a-(16)": "其他侵入竊盜",
    "C-b-(1)": "汽車竊盜", "C-b-(2)": "機車竊盜", "C-b-(3)": "自行車竊盜",
    "C-c-(13)": "搶奪", "C-c-(14)": "扒竊", "C-c-(15)": "置物竊盜", "C-c-(16)": "趁睡竊盜",
    "C-c-(17)": "車內物品竊盜", "C-c-(18)": "零件竊盜", "C-c-(19)": "更衣場所竊盜",
    "C-c-(20)": "自動販賣機竊盜", "C-c-(21)": "內衣等竊盜", "C-c-(22)": "工地竊盜",
    "C-c-(23)": "商店扒竊（萬引き）", "C-c-(24)": "職場竊盜", "C-c-(25)": "同住者竊盜",
    "C-c-(26)": "香油錢竊盜", "C-c-(27)": "其他非侵入竊盜",
    "F-4": "妨害公務", "F-8": "失火", "F-12": "侵入住宅（不含竊盜）",
    "F-27": "非法逮捕監禁", "F-28": "誘拐及人口販運", "F-29": "誹謗",
    "F-32": "侵占遺失物等", "F-35": "建物損壞", "F-37": "器物損壞等",
}
total = counts(books[0]["刑法犯総数"])
definitions = []
data = {name: {"total": total[name], "counts": {}} for name in names}
for code, label in groups:
    parent = counts(books[0][code])
    sheets = ([s for s in books[1] if "-(" in s.title] if code == "C" else
              [s for s in books[0] if s.title.startswith(code + "-") and s.title.count("-") == 1])
    children = []
    child_sums = dict.fromkeys(names, 0)
    for sheet in sheets:
        values = counts(sheet)
        original = re.sub(r"^\S+\s+", "", str(sheet.cell(4, 3).value)).strip()
        children.append({"code": sheet.title, "label": translations.get(sheet.title, original), "original": original})
        for name in names:
            data[name]["counts"][sheet.title] = values[name]
            child_sums[name] += values[name]
    for name in names:
        assert child_sums[name] == parent[name], (name, code, child_sums[name], parent[name])
        data[name]["counts"][code] = parent[name]
    definitions.append({"code": code, "label": label, "items": children})
for name in names:
    assert sum(data[name]["counts"][code] for code, _ in groups) == total[name], name
output = {"year": YEAR, "retrievedAt": datetime.now(timezone.utc).date().isoformat(),
          "sourceUrl": BASE + "R05hanzaitoukei.htm", "files": SOURCES, "groups": definitions, "prefectures": data}
path = ROOT / "src/data/prefectureCrimeCounts.json"
path.write_text(json.dumps(output, ensure_ascii=False, separators=(",", ":")) + "\n")
print(f"Validated 47 prefectures, national sums, six categories and {sum(len(g['items']) for g in definitions)} detail items; wrote {path}")
print("千葉県", data["千葉県"]["total"], {code:data['千葉県']['counts'][code] for code, _ in groups})
